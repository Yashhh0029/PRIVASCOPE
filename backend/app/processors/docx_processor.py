from pathlib import Path
from typing import List
import docx
from app.processors.base import BaseProcessor, ExtractedContent, ContentBlock

class DocxProcessor(BaseProcessor):
    format = "docx"

    def extract(self, file_path: Path) -> ExtractedContent:
        doc = docx.Document(str(file_path))
        blocks: List[ContentBlock] = []
        raw_lines: List[str] = []

        # Extract paragraphs
        for p_idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if text:
                raw_lines.append(text)
                blocks.append(ContentBlock(
                    text=text,
                    page=1,
                    source="docx",
                    row=p_idx + 1,
                    context_meta={"style": p.style.name if p.style else "Normal"}
                ))

        # Extract tables
        for t_idx, table in enumerate(doc.tables):
            header = []
            for r_idx, row in enumerate(table.rows):
                row_texts = [cell.text.strip() for cell in row.cells]
                if r_idx == 0:
                    header = row_texts
                raw_lines.append(" | ".join(row_texts))
                for c_idx, cell_text in enumerate(row_texts):
                    if cell_text:
                        col_name = header[c_idx] if (header and c_idx < len(header)) else f"Col{c_idx+1}"
                        blocks.append(ContentBlock(
                            text=cell_text,
                            page=1,
                            source="docx",
                            row=r_idx + 1,
                            col=c_idx + 1,
                            col_name=col_name,
                            context_meta={"table_index": t_idx}
                        ))

        return ExtractedContent(
            raw_text="\n".join(raw_lines),
            blocks=blocks,
            format=self.format,
            metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}
        )

    def _replace_text_in_paragraph(self, paragraph, val_map: dict[str, str]):
        for orig, repl in val_map.items():
            if orig in paragraph.text:
                # To maintain formatting runs as best as possible
                for run in paragraph.runs:
                    if orig in run.text:
                        run.text = run.text.replace(orig, repl)
                # Fallback if text was split across runs
                if orig in paragraph.text:
                    paragraph.text = paragraph.text.replace(orig, repl)

    def protect(self, file_path: Path, output_path: Path, detections: list, mode: str) -> bool:
        doc = docx.Document(str(file_path))

        val_map = {}
        for d in detections:
            replacement = "[REDACTED]" if mode == "REDACT" else d.masked_preview
            val_map[d.matched_value] = replacement

        for p in doc.paragraphs:
            self._replace_text_in_paragraph(p, val_map)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        self._replace_text_in_paragraph(p, val_map)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return True

    def verify_protection(self, output_path: Path, original_values: list[str]) -> bool:
        if not output_path.exists():
            return False
        try:
            doc = docx.Document(str(output_path))
            all_text = []
            for p in doc.paragraphs:
                all_text.append(p.text)
            for t in doc.tables:
                for row in t.rows:
                    for cell in row.cells:
                        all_text.append(cell.text)
            full_text = " ".join(all_text)
            for val in original_values:
                if val and val in full_text:
                    return False
            return True
        except Exception:
            return False
