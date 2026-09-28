from pathlib import Path
from typing import List
import openpyxl
from app.processors.base import BaseProcessor, ExtractedContent, ContentBlock

class ExcelProcessor(BaseProcessor):
    format = "xlsx"

    def extract(self, file_path: Path) -> ExtractedContent:
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        blocks: List[ContentBlock] = []
        raw_lines: List[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            headers: dict[int, str] = {}

            for row_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
                row_vals = []
                for col_idx, cell in enumerate(row, start=1):
                    val = cell.value
                    if val is None:
                        continue
                    str_val = str(val).strip()
                    if not str_val:
                        continue

                    # If row 1, record as header
                    if row_idx == 1:
                        headers[col_idx] = str_val

                    col_name = headers.get(col_idx, f"Col{col_idx}")
                    row_vals.append(f"{col_name}: {str_val}")

                    blocks.append(ContentBlock(
                        text=str_val,
                        page=1,
                        source="xlsx",
                        sheet_name=sheet_name,
                        row=row_idx,
                        col=col_idx,
                        col_name=col_name,
                        context_meta={"cell_coord": cell.coordinate}
                    ))
                if row_vals:
                    raw_lines.append(f"[{sheet_name}] " + " | ".join(row_vals))

        return ExtractedContent(
            raw_text="\n".join(raw_lines),
            blocks=blocks,
            format=self.format,
            metadata={"sheets": wb.sheetnames}
        )

    def protect(self, file_path: Path, output_path: Path, detections: list, mode: str) -> bool:
        # Load workbook with formatting preserved
        wb = openpyxl.load_workbook(str(file_path))

        val_map = {}
        for d in detections:
            replacement = "[REDACTED]" if mode == "REDACT" else d.masked_preview
            val_map[d.matched_value] = replacement

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    str_val = str(cell.value)
                    modified = False
                    for orig_val, repl in val_map.items():
                        if orig_val in str_val:
                            str_val = str_val.replace(orig_val, repl)
                            modified = True
                    if modified:
                        cell.value = str_val

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        return True

    def verify_protection(self, output_path: Path, original_values: list[str]) -> bool:
        if not output_path.exists():
            return False
        try:
            wb = openpyxl.load_workbook(str(output_path), data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.iter_rows(values_only=True):
                    for cell_val in row:
                        if cell_val is None:
                            continue
                        str_val = str(cell_val)
                        for target in original_values:
                            if target and target in str_val:
                                return False
            return True
        except Exception:
            return False
